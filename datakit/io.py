"""Reading messy input files and writing client-ready Excel reports."""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ENCODINGS = ("utf-8-sig", "cp1250", "iso-8859-2")


def read_csv_smart(path: Path) -> tuple[pd.DataFrame, str]:
    """Read a CSV with unknown encoding and separator (',' ';' tab). Everything as text."""
    for encoding in ENCODINGS:
        try:
            df = pd.read_csv(path, sep=None, engine="python", encoding=encoding, dtype=str)
            return df, encoding
        except UnicodeDecodeError:
            continue
    raise ValueError(f"Could not decode {path} with any of {ENCODINGS}")


def read_excel_smart(path: Path, sheet_name: int | str = 0) -> pd.DataFrame:
    """Read an Excel sheet whose real header is not in row 1 (titles, blank rows above)."""
    raw = pd.read_excel(path, sheet_name=sheet_name, header=None, dtype=str)
    filled = raw.notna().sum(axis=1)
    header_row = int(filled[filled == filled.max()].index[0])
    df = raw.iloc[header_row + 1:].copy()
    df.columns = raw.iloc[header_row].astype(str).str.strip()
    return df.dropna(how="all").reset_index(drop=True)


def write_excel_report(sheets: dict[str, pd.DataFrame], path: Path) -> None:
    """Write several sheets with bold frozen headers, autofilter and fitted column widths."""
    path.parent.mkdir(parents=True, exist_ok=True)
    header_fill = PatternFill("solid", fgColor="E1E0D9")
    with pd.ExcelWriter(path, engine="openpyxl", datetime_format="YYYY-MM-DD", date_format="YYYY-MM-DD") as writer:
        for name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=name[:31], index=False)
            ws = writer.sheets[name[:31]]
            ws.freeze_panes = "A2"
            if frame.shape[1]:
                ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.fill = header_fill
                cell.alignment = Alignment(vertical="center")
            for idx, column in enumerate(frame.columns, start=1):
                lengths = [len(str(v)) for v in frame[column].head(500)] + [len(str(column))]
                ws.column_dimensions[get_column_letter(idx)].width = min(max(lengths) + 2, 60)


@dataclass
class CleaningLog:
    """Audit trail of what was changed — goes into the report delivered to the client."""
    entries: list[tuple[str, int]] = field(default_factory=list)

    def add(self, step: str, count: int) -> None:
        self.entries.append((step, int(count)))
        print(f"  {count:>6}  {step}")

    def to_frame(self) -> pd.DataFrame:
        return pd.DataFrame(self.entries, columns=["krok", "liczba"])

    def to_markdown(self) -> str:
        rows = "\n".join(f"| {step} | {count} |" for step, count in self.entries)
        return "| Krok | Liczba |\n|---|---:|\n" + rows
